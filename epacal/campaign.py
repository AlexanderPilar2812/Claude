"""Análisis y depuración de una campaña de medición (presión, caudal, nivel) con el INP.

Entradas: el INP del modelo, el Excel de la campaña (una hoja por equipo) y los puntos
de medición en GeoPackage. Salidas:

* qué pasó probablemente en la campaña (cortes, cambios de operación, VRP que no regulan,
  sensores mal ubicados, tanques que rebosan, sensores fuera del agua…);
* para cada sensor, qué datos son válidos y cuáles se reemplazan, y con qué método;
* un Excel depurado con la marca de cada dato, ficheros .dat de día tipo para EPACAL
  y un informe HTML.

Métodos de reemplazo (columna ``marca``):

    0  válido                    dato medido, sin tocar
    1  envolvente                oscilación más rápida que el registro: media recortada 2 h
    2  interpolado               pico aislado o hueco corto (≤ 1 h)
    3  perfil típico             hueco medio (≤ 6 h): perfil diario del propio sensor,
                                 ajustado a los extremos del hueco
    4  balance de masa           nivel de tanque reconstruido con caudales de entrada/salida
    9  excluido                  fuera de servicio, corte de suministro, sensor fuera del agua
"""
from __future__ import annotations

import collections
import datetime as dt
import html
import io
import os
import re
import sqlite3
import struct
import base64
from dataclasses import dataclass, field
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from . import engine as en
from .network import parse_coordinates

STEP = "15min"
FLAG_TXT = {0: "válido", 1: "envolvente (oscilación rápida)", 2: "interpolado",
            3: "perfil típico", 4: "balance de masa", 9: "excluido"}
KIND_TXT = {"PRESSURE": "Presión", "FLOW": "Caudal", "LEVEL": "Nivel"}
UNIT = {"PRESSURE": "m", "FLOW": "L/s", "LEVEL": "m"}
# factor para pasar de las unidades de caudal del INP a L/s
TO_LPS = {"LPS": 1.0, "LPM": 1 / 60, "MLD": 1000 / 86.4, "CMH": 1 / 3.6, "CMD": 1 / 86.4,
          "CFS": 28.3168, "GPM": 0.0630902, "MGD": 43.8126, "IMGD": 52.6168, "AFD": 14.2764}


# --------------------------------------------------------------------------- datos
@dataclass
class Sensor:
    name: str
    kind: str                       # PRESSURE | FLOW | LEVEL
    raw: pd.Series                  # valores en m o L/s, con su hora real
    title: str = ""
    point: Optional[str] = None     # nombre del punto en el GeoPackage
    xy: Optional[Tuple[float, float]] = None
    element: Optional[str] = None   # nudo / línea / tanque del INP
    elev: float = float("nan")
    elev_meas: float = float("nan")      # cota del sensor medida (GNSS) si viene en el GPKG
    zone: Optional[int] = None
    prv: Optional[str] = None       # VRP que alimenta su zona
    notes: List[Tuple[str, str]] = field(default_factory=list)   # (nivel, texto)
    # resultados
    grid: Optional[pd.Series] = None
    clean: Optional[pd.Series] = None
    flag: Optional[pd.Series] = None
    offset: float = 0.0
    t_on: Optional[pd.Timestamp] = None
    t_off: Optional[pd.Timestamp] = None
    verdict: str = ""
    model_day: Optional[np.ndarray] = None
    weight: float = 1.0

    def note(self, level: str, text: str) -> None:
        self.notes.append((level, text))


# --------------------------------------------------------------------------- lectura
def read_gpkg_points(path: str) -> Dict[str, dict]:
    """Puntos de un GeoPackage: {Nombre: {x, y, campos…}} (sin depender de GDAL)."""
    con = sqlite3.connect(path)
    out = {}
    for (table,) in con.execute("select table_name from gpkg_contents where data_type='features'"):
        gcol = con.execute("select column_name from gpkg_geometry_columns where table_name=?",
                           (table,)).fetchone()[0]
        cur = con.execute(f'select * from "{table}"')
        cols = [d[0] for d in cur.description]
        for row in cur.fetchall():
            rec = dict(zip(cols, row))
            b = rec.pop(gcol)
            if not b:
                continue
            env = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}[(b[3] >> 1) & 7]
            w = b[8 + env:]
            x, y = struct.unpack(("<" if w[0] == 1 else ">") + "dd", w[5:21])
            name = str(rec.get("Nombre") or rec.get("nombre") or rec.get("name") or rec.get("fid"))
            text = " ".join(str(v) for v in rec.values() if isinstance(v, str))
            out[name] = {"x": x, "y": y, "text": text, "layer": table, "attrs": rec}
    con.close()
    return out


def _kind_of(header: str) -> Optional[str]:
    h = header.lower()
    if "l/s" in h or "caudal" in h:
        return "FLOW"
    if "nivel" in h:
        return "LEVEL"
    if "mca" in h or "presi" in h:
        return "PRESSURE"
    return None


def read_campaign_excel(path: str, log: Callable[[str], None] = print,
                        sheets: Optional[Dict[str, dict]] = None) -> List[Sensor]:
    """Una hoja por equipo: fila de cabecera que empieza por «Fecha», columna de valores en
    m / MCA / L/s. Se ignoran las hojas de resumen (varias columnas de sensores)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        spec = (sheets or {}).get(ws.title, {})
        if spec.get("skip"):
            continue
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = next((i for i, r in enumerate(rows[:15])
                      if r and isinstance(r[0], str) and r[0].strip().lower().startswith("fecha")), None)
        if hdr_i is None:
            continue
        hdr = [str(h or "").replace("\n", " ") for h in rows[hdr_i]]
        if "col" in spec:
            vcol = int(spec["col"])
        else:
            vcol = None
            for j, h in enumerate(hdr[1:6], start=1):
                hl = h.lower()
                if any(w in hl for w in ("offset", "depurad", "prom", "cmh", "mbar")):
                    continue
                if ("(mca)" in hl or "(m)" in hl or "l/s" in hl) and _kind_of(h):
                    vcol = j
                    break
            # hojas consolidadas: muchas columnas de sensores → no son de un equipo
            if vcol is None or sum("(mca)" in h.lower() for h in hdr) > 2:
                continue
        kind = spec.get("type") or _kind_of(hdr[vcol]) or "PRESSURE"
        t, v = [], []
        for r in rows[hdr_i + 1:]:
            if len(r) > vcol and isinstance(r[0], dt.datetime) and isinstance(r[vcol], (int, float)):
                t.append(r[0])
                v.append(float(r[vcol]))
        if len(t) < 10:
            continue
        s = pd.Series(v, index=pd.DatetimeIndex(t)).sort_index()
        s = s[~s.index.duplicated()]
        title = " ".join(str(c) for r in rows[:hdr_i] for c in r if isinstance(c, str))[:200]
        name = spec.get("name") or ws.title
        if kind == "LEVEL":
            m = re.search(r"TANQUE\s+([A-Z0-9ÁÉÍÓÚÑ]+)", title.upper())
            if m and not spec.get("name"):
                name = "N_" + m.group(1).capitalize()
        out.append(Sensor(name, kind, s, title=title, point=spec.get("point")))
        log(f"  {ws.title}: {KIND_TXT[kind]} «{name}», {len(s)} registros "
            f"{s.index[0]:%d/%m %H:%M} → {s.index[-1]:%d/%m %H:%M}")
    return out


# --------------------------------------------------------------------------- red
class Net:
    """Lo que hace falta del INP: coordenadas, cotas, zonas de presión y resultados del modelo."""

    def __init__(self, inp: str, lib: Optional[str] = None):
        self.inp = inp
        self.p = en.EpanetProject(inp, lib or en._find_library(None, "2.2"))
        p = self.p
        self.n = p.count(en.EN_NODECOUNT)
        self.L = p.count(en.EN_LINKCOUNT)
        self.nid = [p.node_id(i) for i in range(1, self.n + 1)]
        self.ntype = np.array([p.node_type(i) for i in range(1, self.n + 1)])
        self.elev = np.array([p.get_node(i, en.EN_ELEVATION) for i in range(1, self.n + 1)])
        coords, _ = parse_coordinates(inp)
        self.xy = np.array([coords.get(k, (np.nan, np.nan)) for k in self.nid], float)
        self.lid = [p.link_id(l) for l in range(1, self.L + 1)]
        self.ltype = np.array([p.link_type(l) for l in range(1, self.L + 1)])
        self.lnodes = np.array([p.link_nodes(l) for l in range(1, self.L + 1)]) - 1
        self.units = p.flow_units()
        self.q2lps = TO_LPS.get(self.units.upper(), 1.0)
        self._zones()
        self.tank_diam = {self.nid[i]: p.get_node(i + 1, en.EN_TANKDIAM)
                          for i in range(self.n) if self.ntype[i] == en.EN_TANK}
        self.tank_lim = {self.nid[i]: (p.get_node(i + 1, en.EN_MINLEVEL), p.get_node(i + 1, en.EN_MAXLEVEL))
                         for i in range(self.n) if self.ntype[i] == en.EN_TANK}

    def _zones(self):
        """Zonas de presión: componentes conexas sin cruzar VRP/VSP/VCQ, bombas ni líneas cerradas."""
        p = self.p
        adj = collections.defaultdict(list)
        self.prv = []
        for l in range(self.L):
            a, b = self.lnodes[l]
            t = self.ltype[l]
            if t in (en.EN_PRV, en.EN_PSV, en.EN_FCV, en.EN_PUMP):
                if t == en.EN_PRV:
                    self.prv.append((self.lid[l], a, b, p.get_link(l + 1, en.EN_INITSETTING)))
                continue
            if p.get_link(l + 1, en.EN_INITSTATUS) == 0:
                continue
            adj[a].append(b)
            adj[b].append(a)
        zone = -np.ones(self.n, int)
        z = 0
        for i in range(self.n):
            if zone[i] >= 0:
                continue
            stack, zone[i] = [i], z
            while stack:
                u = stack.pop()
                for v in adj[u]:
                    if zone[v] < 0:
                        zone[v] = z
                        stack.append(v)
            z += 1
        self.zone = zone
        self.zone_prv = {}
        for k, a, b, s in self.prv:
            self.zone_prv.setdefault(zone[b], []).append(
                {"id": k, "up_zone": zone[a], "setting": s, "target": self.elev[b] + s, "up": a, "down": b})

    def mean_head(self, nodes) -> np.ndarray:
        """Carga media en el modelo (simulación completa, guardada en caché)."""
        if getattr(self, "_mean_head", None) is None:
            self._mean_head = self.simulate(list(range(self.n)), [])["H"].mean(0)
        return self._mean_head[np.asarray(nodes)]

    def nearest_node(self, x, y, types=None) -> Tuple[int, float]:
        d = np.hypot(self.xy[:, 0] - x, self.xy[:, 1] - y)
        if types is not None:
            d = np.where(np.isin(self.ntype, types), d, np.inf)
        i = int(np.nanargmin(d))
        return i, float(d[i])

    def nearest_link(self, x, y) -> Tuple[int, float]:
        a, b = self.xy[self.lnodes[:, 0]], self.xy[self.lnodes[:, 1]]
        ab = b - a
        t = np.clip(((x - a[:, 0]) * ab[:, 0] + (y - a[:, 1]) * ab[:, 1]) /
                    np.maximum((ab ** 2).sum(1), 1e-9), 0, 1)
        d = np.hypot(a[:, 0] + t * ab[:, 0] - x, a[:, 1] + t * ab[:, 1] - y)
        d = np.where(self.ltype == en.EN_PIPE, d, d + 5.0)   # preferir tuberías
        i = int(np.nanargmin(d))
        return i, float(d[i])

    def simulate(self, nodes: Sequence[int], links: Sequence[int]) -> dict:
        """Presión/carga en nudos y caudal en líneas, cada 15 min, durante la duración del INP."""
        p = self.p
        nodes = [int(i) for i in nodes]
        links = [int(l) for l in links]
        rec = {"t": [], "P": [], "H": [], "Q": []}

        def cb(t):
            rec["t"].append(t)
            rec["P"].append([p.get_node(i + 1, en.EN_PRESSURE) for i in nodes])
            rec["H"].append([p.get_node(i + 1, en.EN_HEAD) for i in nodes])
            rec["Q"].append([p.get_link(l + 1, en.EN_FLOW) * self.q2lps for l in links])
        p.run_eps(900, cb)
        return {k: np.array(v) for k, v in rec.items()}


def _day_profile_from_sim(t: np.ndarray, v: np.ndarray) -> np.ndarray:
    """Media por cuarto de hora del día (descarta las primeras 24 h si hay más de 2 días)."""
    slot = ((t // 900) % 96).astype(int)
    keep = t >= 86400 if t.max() >= 2 * 86400 else np.ones_like(t, bool)
    prof = np.full(96, np.nan)
    for k in range(96):
        m = keep & (slot == k)
        if m.any():
            prof[k] = v[m].mean()
    return pd.Series(prof).interpolate(limit_direction="both").to_numpy()


# --------------------------------------------------------------------------- ubicación
def locate(sensors: List[Sensor], points: Dict[str, dict], net: Net, log=print) -> None:
    """Asocia cada sensor a su punto del GPKG y a su elemento del INP."""
    flowpts = {k: v for k, v in points.items() if "caudal" in v["layer"].lower() or k.upper().startswith("C")}
    for s in sensors:
        pt = s.point
        if pt is None:
            cands = [s.name, s.name.upper()]
            if s.kind == "FLOW" and re.match(r"^Q\d+$", s.name):
                cands.append("C" + s.name[1:])
            pt = next((c for c in cands if c in points), None)
        if pt is None and s.kind == "LEVEL":
            # el nombre del tanque aparece en la descripción de un caudalímetro cercano
            tname = s.name.split("_", 1)[-1].lower()
            for k, v in points.items():
                if not (k.upper().startswith("N") or "nivel" in v["layer"].lower()):
                    continue
                near = [f for f in flowpts.values() if np.hypot(f["x"] - v["x"], f["y"] - v["y"]) < 60]
                if any(tname in f["text"].lower() for f in near):
                    pt = k
                    break
            if pt is None:
                # el nombre del tanque aparece en el título de un caudal medido cerca (≤ 250 m)
                for f in sensors:
                    if f.kind != "FLOW" or tname not in f.title.lower():
                        continue
                    fp = (f.point or ("C" + f.name[1:] if re.match(r"^Q\d+$", f.name) else None))
                    if fp not in points:
                        continue
                    lv = [(np.hypot(points[fp]["x"] - v["x"], points[fp]["y"] - v["y"]), k)
                          for k, v in points.items() if k.upper().startswith("N") or "nivel" in v["layer"].lower()]
                    lv = [x for x in lv if x[0] < 250]
                    if lv:
                        pt = min(lv)[1]
                        break
        s.point = pt
        if pt is None:
            s.note("error", "No tiene punto en el GeoPackage: no se puede ubicar en el modelo. "
                            "Añádalo a la capa de puntos antes de usarlo en la calibración.")
            continue
        x, y = points[pt]["x"], points[pt]["y"]
        s.xy = (x, y)
        at = points[pt].get("attrs", {})
        for k in ("Elevacion_Medida", "elev_gnss", "Elevacion_msnm"):
            if isinstance(at.get(k), (int, float)):
                s.elev_meas = float(at[k])
                break
        if s.kind == "FLOW":
            l, d = net.nearest_link(x, y)
            s.element = net.lid[l]
            s._link = l
        elif s.kind == "LEVEL":
            i, d = net.nearest_node(x, y, [en.EN_TANK, en.EN_RESERVOIR])
            s.element, s._node = net.nid[i], i
            s.elev = net.elev[i]
            if net.ntype[i] == en.EN_RESERVOIR:
                s.reservoir = True
                s.note("info", f"En el modelo es un embalse de carga fija ({net.elev[i]:.1f} m). Su nivel medido "
                               f"no es una meta de calibración: úselo para el patrón de carga del embalse "
                               f"(cota de fondo + nivel) o conviértalo en tanque.")
        else:
            i, d = net.nearest_node(x, y, [en.EN_JUNCTION])
            dd = np.hypot(net.xy[:, 0] - x, net.xy[:, 1] - y)
            junc = net.ntype == en.EN_JUNCTION
            station = np.where(junc & (dd <= d + 3.0))[0]
            zs = {int(net.zone[k]) for k in station}
            p_meas = _pressure_level(s)
            if len(zs) > 1 and p_meas is not None:
                # estación de válvulas (varios nudos en el mismo sitio, en zonas distintas):
                # el nudo cuya carga en el modelo se parece más a la medida
                h = net.mean_head(station)
                i = int(station[np.argmin(np.abs(p_meas + net.elev[station] - h))])
                s.note("info", f"El punto está en una estación de válvulas ({len(station)} nudos en el mismo "
                               f"sitio): se asigna al nudo {net.nid[i]}, cuya carga en el modelo coincide con "
                               f"la medida.")
            elif np.isfinite(s.elev_meas):
                # con cota GNSS: entre los nudos cercanos, el de cota más parecida
                near = np.where((dd < max(30.0, 1.5 * d)) & junc)[0]
                ok = near[np.abs(net.elev[near] - s.elev_meas) < 10]
                if len(ok):
                    i = int(ok[np.argmin(dd[ok] + 3.0 * np.abs(net.elev[ok] - s.elev_meas))])
            d = float(dd[i])
            s.element, s._node = net.nid[i], i
            s.elev = net.elev[i]
            s.zone = int(net.zone[i])
            if np.isfinite(s.elev_meas) and abs(s.elev_meas - s.elev) > 10:
                s.note("aviso", f"La cota GNSS del sensor ({s.elev_meas:.2f} m) no es compatible con el sitio: "
                                f"el nudo {s.element} está a {s.elev:.2f} m y ningún nudo cercano tiene esa cota. "
                                f"Con la cota del nudo la carga medida es {(_pressure_level(s) or 0) + s.elev:.1f} m. "
                                f"Revise la cota GNSS (¿es la de otro punto?); no se usa para corregir la presión.")
                s.elev_meas = float("nan")
            elif np.isfinite(s.elev_meas) and abs(s.elev_meas - s.elev) > 2:
                s.note("aviso", f"Cota GNSS del sensor {s.elev_meas:.2f} m y cota del nudo {s.element} en el INP "
                                f"{s.elev:.2f} m (Δ {s.elev_meas - s.elev:+.2f} m). La presión del modelo se compara "
                                f"pasando la medida a la cota del nudo; revise la cota del nudo.")
            fe = net.zone_prv.get(s.zone, [])
            s.prv = fe[0]["id"] if fe else None
        if d > 25:
            s.note("aviso", f"El punto {pt} está a {d:.0f} m del elemento más cercano del modelo "
                            f"({s.element}); compruebe la ubicación.")
        log(f"  {s.name:12s} → punto {pt:4s} → {s.element} (a {d:.1f} m)")


# --------------------------------------------------------------------------- utilidades
def _pressure_level(s: "Sensor") -> Optional[float]:
    """Presión típica medida (sin offset, con el equipo instalado) para ubicar el sensor."""
    if s.clean is None:
        return None
    x = s.clean.where(s.flag.isin([0, 1, 2, 3])).dropna()
    return float(x.median()) if len(x) > 20 else None


def to_grid(s: pd.Series, step: str = STEP, max_gap: str = "40min") -> pd.Series:
    """Remuestrea a una malla regular interpolando en el tiempo real de cada lectura
    (los equipos no registran a la misma hora en punto)."""
    s = s.dropna()
    if pd.Timedelta(step) < s.index.to_series().diff().median() * 1.5:
        g = pd.date_range(s.index[0].ceil(step), s.index[-1].floor(step), freq=step)
        x = s.index.asi8.astype(float)
        y = np.interp(g.asi8.astype(float), x, s.to_numpy())
        # no inventar dentro de huecos largos
        pos = np.searchsorted(x, g.asi8.astype(float))
        pos = np.clip(pos, 1, len(x) - 1)
        gap = (x[pos] - x[pos - 1]) / 1e9
        y[gap > pd.Timedelta(max_gap).total_seconds()] = np.nan
        return pd.Series(y, index=g)
    return s.resample(step).mean()


def trimmed_roll(x: pd.Series, win: int = 9, trim: float = 0.25) -> pd.Series:
    def tm(a):
        a = np.sort(a[np.isfinite(a)])
        if len(a) == 0:
            return np.nan
        k = int(len(a) * trim)
        return a[k:len(a) - k].mean() if len(a) > 2 * k else a.mean()
    return x.rolling(win, center=True, min_periods=max(3, win // 2)).apply(tm, raw=True)


def runs(mask: np.ndarray) -> List[Tuple[int, int]]:
    """Tramos [i, j) donde mask es True."""
    m = np.concatenate([[False], np.asarray(mask, bool), [False]])
    d = np.diff(m.astype(int))
    return list(zip(np.where(d == 1)[0], np.where(d == -1)[0]))


def fmt(t) -> str:
    return pd.Timestamp(t).strftime("%d/%m %H:%M")


# --------------------------------------------------------------------------- diagnóstico por sensor
def detect_service(s: Sensor, atm: float = 3.0) -> None:
    """Instalación y retiro (meseta atmosférica en los extremos) y offset de cero."""
    g = s.grid
    if s.kind == "PRESSURE":
        on = (g > atm).to_numpy()
        if on.any():
            i0, i1 = int(np.argmax(on)), len(on) - int(np.argmax(on[::-1]))
            pre = g.iloc[:i0]
            post = g.iloc[i1:]
            base = pd.concat([pre, post])
            base = base[base.abs() < atm]
            if len(base) >= 3:
                # moda de la meseta (descarta el transitorio de puesta en marcha)
                vals = s.raw[(s.raw.abs() < atm)]
                vals = vals[(vals.index < g.index[i0]) | (vals.index > g.index[i1 - 1])]
                if len(vals) >= 3:
                    r = vals.round(2)
                    s.offset = float(r.mode().iloc[0]) if r.duplicated().any() else float(vals.median())
            # quitar también la lectura de transición (sensor a medio presurizar)
            s.t_on, s.t_off = g.index[min(i0 + 1, len(g) - 1)], g.index[max(i1 - 2, 0)]
    elif s.kind == "LEVEL":
        s.t_on, s.t_off = g.index[0], g.index[-1]
    else:
        s.t_on, s.t_off = g.index[0], g.index[-1]
    if s.offset:
        s.note("info", f"Offset de cero (lectura a presión atmosférica antes/después de instalar): "
                       f"{s.offset:+.2f} m; se resta a toda la serie.")


def daily_noise(x: pd.Series) -> pd.Series:
    return x.diff().abs().groupby(x.index.normalize()).median()


def analyse_sensor(s: Sensor, ctx: dict) -> None:
    g = (s.grid - s.offset) if s.kind == "PRESSURE" else s.grid.copy()
    flag = pd.Series(0, index=g.index, dtype=int)
    reason = pd.Series("", index=g.index, dtype=object)
    clean = g.copy()

    outside = (g.index < s.t_on) | (g.index > s.t_off)
    flag[outside] = 9
    reason[outside] = "equipo sin instalar"

    if s.kind == "LEVEL":
        out_water = _level_out_of_water(s, g)
        if out_water is not None:
            flag[out_water] = 9
            reason[out_water] = "sensor fuera del agua"

    valid = (flag == 0) & g.notna()
    base = g.where(valid).rolling(9, center=True, min_periods=4).median()
    resid = g - base

    # ---- caídas a cero dentro del servicio (cortes / vaciados)
    if s.kind == "PRESSURE":
        low = valid & (g < 3.0)
        for i, j in runs(low.to_numpy()):
            flag.iloc[i:j] = 9
            reason.iloc[i:j] = "sin presión (corte o vaciado)"
        valid = (flag == 0) & g.notna()

    # ---- congelado (mismo valor ≥ 2 h); un caudal nulo es una válvula cerrada, no un fallo
    same = g.diff().abs() < 1e-9
    for i, j in runs(same.to_numpy() & valid.to_numpy()):
        if j - i >= 8 and not (s.kind == "FLOW" and abs(g.iloc[i]) < 1e-6):
            flag.iloc[i - 1:j] = 9
            reason.iloc[i - 1:j] = "sensor congelado (valor repetido)"
    valid = (flag == 0) & g.notna()

    # ---- oscilación más rápida que el registro: días con ruido muy superior al habitual
    nd = daily_noise(g.where(valid))
    cover = g.where(valid).notna().groupby(g.index.normalize()).mean()
    nd = nd[cover > 0.5]
    fast_days = []
    if len(nd) >= 2 and s.kind != "FLOW":
        # un salto mediano entre lecturas mayor que el umbral no lo reproduce un modelo EPS
        for d, v in nd.items():
            if v > ctx["noise_abs"][s.kind]:
                fast_days.append(d)
    s._fast_days = fast_days
    s._noise = nd
    if fast_days:
        strong = float(nd[nd.index.isin(fast_days)].median()) > 2 * ctx["noise_abs"][s.kind]
        env = trimmed_roll(g.where(valid), 17 if strong else 9)
        m = valid & g.index.normalize().isin(fast_days)
        clean[m] = env[m]
        flag[m] = 1
        reason[m] = f"oscilación más rápida que el registro → media recortada de {4 if strong else 2} h"

    # ---- picos aislados (Hampel) en el resto
    q = valid & (flag == 0)
    sig = 1.4826 * (resid[q]).abs().median()
    thr = max(4 * sig, ctx["spike_abs"][s.kind])
    spike = q & (resid.abs() > thr)
    for i, j in runs(spike.to_numpy()):
        if j - i <= 3:
            flag.iloc[i:j] = 2
            reason.iloc[i:j] = "pico aislado"
            clean.iloc[i:j] = np.nan

    # ---- huecos: cortos → interpolación; medios → perfil típico; largos → excluido
    clean[flag == 9] = np.nan
    gaps = clean.isna() & (flag != 9)
    gaps |= g.isna() & ~pd.Series(outside, index=g.index)
    prof = _typical_profile(clean.where(flag == 0))
    for i, j in runs(gaps.to_numpy()):
        n = j - i
        if n <= 4:
            seg = clean.iloc[max(i - 1, 0):min(j + 1, len(clean))].interpolate(limit_area="inside")
            clean.iloc[i:j] = seg.iloc[(1 if i > 0 else 0):(1 if i > 0 else 0) + n].to_numpy()
            flag.iloc[i:j] = np.where(flag.iloc[i:j] == 0, 2, flag.iloc[i:j])
            reason.iloc[i:j] = np.where(reason.iloc[i:j] == "", "hueco corto", reason.iloc[i:j])
        elif n <= 24 and prof is not None and i > 0 and j < len(clean):
            slots = ((clean.index[i:j].hour * 60 + clean.index[i:j].minute) // 15).to_numpy()
            pr = prof[slots]
            a0 = clean.iloc[i - 1] - prof[(clean.index[i - 1].hour * 60 + clean.index[i - 1].minute) // 15]
            a1 = clean.iloc[j] - prof[(clean.index[j].hour * 60 + clean.index[j].minute) // 15]
            w = np.linspace(0, 1, n + 2)[1:-1]
            clean.iloc[i:j] = pr + (1 - w) * a0 + w * a1
            flag.iloc[i:j] = 3
            reason.iloc[i:j] = "hueco de " + f"{n / 4:.1f} h → perfil típico ajustado a los extremos"
    if s.kind == "FLOW":
        x = g[valid]
        zero = (x.abs() < 0.05).mean()
        on = x[x.abs() > 0.05]
        s.onoff = bool(0.15 < zero < 0.95 and len(on) > 10 and on.std() < 0.35 * on.mean())
    s.clean, s.flag, s.reason = clean, flag, reason


def _typical_profile(x: pd.Series, mean: bool = False) -> Optional[np.ndarray]:
    """Día tipo por cuarto de hora: mediana (media en caudales todo/nada)."""
    x = x.dropna()
    if len(x) < 96:
        return None
    slot = (x.index.hour * 60 + x.index.minute) // 15
    gb = x.groupby(slot)
    prof = (gb.mean() if mean else gb.median()).reindex(range(96))
    return prof.interpolate(limit_direction="both").to_numpy()


def _level_out_of_water(s: Sensor, g: pd.Series) -> Optional[np.ndarray]:
    """Un nivel negativo o casi constante durante días es un sensor fuera del agua."""
    bad = (g < -0.05).to_numpy()
    roll_std = g.rolling(48, center=True, min_periods=12).std().to_numpy()
    flat = roll_std < 0.01
    m = bad | (flat & (g.abs() < 0.3).to_numpy())
    if m.mean() > 0.05:
        frac = m.mean() * 100
        s.note("error", f"Sensor de nivel fuera del agua el {frac:.0f} % del registro "
                        f"(lectura constante {np.nanmedian(g[m]):.2f} m). Esos datos no son nivel.")
        return m
    return None


# --------------------------------------------------------------------------- análisis global
def find_events(sensors: List[Sensor]) -> List[dict]:
    """Cortes simultáneos: varios sensores de presión caen a casi cero a la vez."""
    ps = [s for s in sensors if s.kind == "PRESSURE" and s.flag is not None]
    if len(ps) < 2:
        return []
    drop = pd.DataFrame({s.name: ((s.grid - s.offset) < 0.6 * (s.grid - s.offset).rolling(9, center=True,
                                  min_periods=3).median()) & (s.flag != 9) | (s.reason == "sin presión (corte o vaciado)")
                         for s in ps}).fillna(False)
    n = drop.sum(1)
    events = []
    for i, j in runs((n >= 3).to_numpy()):
        idx = drop.index[i:j]
        who = [c for c in drop.columns if drop[c].iloc[i:j].any()]
        events.append({"inicio": idx[0], "fin": idx[-1] + pd.Timedelta(STEP), "tipo": "corte",
                       "sensores": who,
                       "texto": f"Caída simultánea de presión en {len(who)} sensores ({', '.join(who)}): "
                                f"maniobra, corte de suministro o vaciado de la red."})
    for s in ps:
        for i, j in runs((s.reason == "sin presión (corte o vaciado)").to_numpy()):
            t0 = s.grid.index[i]
            if not any(e["inicio"] - pd.Timedelta("1h") <= t0 <= e["fin"] for e in events):
                events.append({"inicio": t0, "fin": s.grid.index[j - 1] + pd.Timedelta(STEP), "tipo": "sin presión",
                               "sensores": [s.name],
                               "texto": f"{s.name} se queda sin presión {(j - i) / 4:.1f} h sin que los demás "
                                        f"lo noten: vaciado local, válvula cerrada aguas arriba o sensor."})
    for s in sensors:
        if s.clean is None:
            continue
        ok = (s.flag != 9) & s.clean.notna()
        if s.kind == "LEVEL":
            for i, j in runs((ok & (s.clean < 0.15)).to_numpy()):
                if j - i >= 4:
                    events.append({"inicio": s.clean.index[i], "fin": s.clean.index[j - 1] + pd.Timedelta(STEP),
                                   "tipo": "tanque vacío", "sensores": [s.name],
                                   "texto": f"{s.name}: nivel por debajo de 0,15 m durante {(j - i) / 4:.1f} h "
                                            f"(tanque vacío: la demanda supera a la entrada)."})
        if s.kind == "FLOW" and not getattr(s, "onoff", False):
            med = max(s.clean[ok].median(), 0.1)
            low = (ok & (s.clean.abs() < 0.3 * med)).astype(float).rolling(5, center=True, min_periods=1).mean() > 0.5
            for i, j in runs(low.to_numpy()):
                if j - i >= 4:
                    events.append({"inicio": s.clean.index[i], "fin": s.clean.index[j - 1] + pd.Timedelta(STEP),
                                   "tipo": "caudal nulo", "sensores": [s.name],
                                   "texto": f"{s.name}: caudal casi nulo o muy reducido (< 30 % del habitual) durante "
                                            f"{(j - i) / 4:.1f} h (válvula cerrada o maniobrada)."})
    events = sorted(events, key=lambda e: e["inicio"])
    merged = []
    for e in events:
        prev = next((m for m in reversed(merged) if m["tipo"] == e["tipo"] and m["sensores"] == e["sensores"]), None)
        if prev and e["tipo"] in ("tanque vacío", "caudal nulo", "sin presión") and \
                e["inicio"] - prev["fin"] <= pd.Timedelta("2h"):
            prev["fin"] = max(prev["fin"], e["fin"])
            h = (prev["fin"] - prev["inicio"]).total_seconds() / 3600
            prev["texto"] = re.sub(r"durante [\d.]+ h", f"durante {h:.1f} h", prev["texto"])
            continue
        merged.append(e)
    return merged


def find_regimes(sensors: List[Sensor], min_days: int = 1) -> List[dict]:
    """Días con el mismo comportamiento: segmentación binaria de los días según
    media, amplitud diaria y ruido de cada sensor de presión/nivel/caudal."""
    feats = {}
    for s in sensors:
        if s.clean is None:
            continue
        x = s.clean.where(s.flag.isin([0, 1, 2, 3]))
        g = s.grid.where(s.flag != 9) - (s.offset if s.kind == "PRESSURE" else 0)
        day = x.index.normalize()
        cov = x.notna().groupby(day).mean()
        mean = x.groupby(day).median()
        amp = x.groupby(day).quantile(0.95) - x.groupby(day).quantile(0.05)
        noi = g.diff().abs().groupby(g.index.normalize()).median()
        for nm, v in (("med", mean), ("amp", amp), ("ruido", noi)):
            v = v.where(cov > 0.6)
            sc = np.nanmedian(np.abs(v - np.nanmedian(v))) * 1.4826 + {"PRESSURE": 1.0, "FLOW": 0.3,
                                                                     "LEVEL": 0.1}[s.kind]
            feats[f"{s.name}:{nm}"] = (v - np.nanmedian(v)) / sc
    F = pd.DataFrame(feats)
    F = F[F.notna().mean(1) > 0.6]
    if len(F) < 2:
        return []
    F = F.loc[:, F.notna().mean() > 0.7]
    F = F.apply(lambda c: c.fillna(c.median()))
    X = F.to_numpy()
    days = list(F.index)

    def cost(a, b):
        return ((X[a:b] - X[a:b].mean(0)) ** 2).sum()

    def split(a, b, out):
        if b - a < 2 * min_days:
            return
        tot = cost(a, b)
        best, k = 0, None
        for m in range(a + min_days, b - min_days + 1):
            gain = tot - cost(a, m) - cost(m, b)
            if gain > best:
                best, k = gain, m
        # ganancia significativa frente a la variabilidad dentro de los tramos
        if k is not None and best > 0.25 * tot and best > 2.0 * X.shape[1]:
            out.append(k)
            split(a, k, out)
            split(k, b, out)

    cuts = []
    split(0, len(days), cuts)
    cuts = sorted(cuts)
    bounds = [0] + cuts + [len(days)]
    regs = []
    for a, b in zip(bounds[:-1], bounds[1:]):
        regs.append({"desde": days[a], "hasta": days[b - 1] + pd.Timedelta("1D"), "dias": b - a,
                     "X": X[a:b].mean(0)})
    # afinar cada frontera: el salto mayor del nivel medio común en ±1 día
    for r0, r1 in zip(regs[:-1], regs[1:]):
        t = _refine_boundary(sensors, r1["desde"])
        if t is not None:
            r0["hasta"] = r1["desde"] = t
    for r in regs:
        r.pop("X")
    return regs


def _refine_boundary(sensors, day) -> Optional[pd.Timestamp]:
    lo, hi = day - pd.Timedelta("18h"), day + pd.Timedelta("30h")
    votes = []
    for s in sensors:
        if s.kind != "PRESSURE" or s.clean is None:
            continue
        x = s.clean[lo:hi]
        if x.notna().mean() < 0.6:
            continue
        a = x.rolling(16, min_periods=8).median()
        b = x[::-1].rolling(16, min_periods=8).median()[::-1]
        jump = (b.shift(-1) - a).abs()
        sc = x.diff().abs().median() + 1e-6
        if jump.max() / sc > 6:
            votes.append(jump.idxmax())
    if len(votes) < 2:
        return None
    v = pd.Series(votes).sort_values()
    # el instante con más votos en ±2 h
    best = max(v, key=lambda t: ((v - t).abs() <= pd.Timedelta("2h")).sum())
    near = v[(v - best).abs() <= pd.Timedelta("2h")]
    ns = near.to_numpy().astype("datetime64[ns]").astype("int64")
    return pd.Timestamp(int(np.median(ns))).floor(STEP) + pd.Timedelta(STEP)


# --------------------------------------------------------------------------- comprobación con el modelo
def model_checks(sensors: List[Sensor], net: Net, window: Tuple[pd.Timestamp, pd.Timestamp],
                 log=print) -> dict:
    nodes = [s._node for s in sensors if s.kind in ("PRESSURE", "LEVEL") and hasattr(s, "_node")]
    links = [s._link for s in sensors if s.kind == "FLOW" and hasattr(s, "_link")]
    up_nodes = sorted({z["up"] for zs in net.zone_prv.values() for z in zs})
    sim = net.simulate(nodes + up_nodes, links)
    t = sim["t"]
    out = {"sim_ok": True}
    k = 0
    a, b = window
    for s in sensors:
        if s.kind in ("PRESSURE", "LEVEL") and hasattr(s, "_node"):
            if s.kind == "PRESSURE":
                s.model_day = _day_profile_from_sim(t, sim["P"][:, k])
            else:
                s.model_day = _day_profile_from_sim(t, sim["H"][:, k] - net.elev[s._node])
            k += 1
    up_head = {n: _day_profile_from_sim(t, sim["H"][:, len(nodes) + j]) for j, n in enumerate(up_nodes)}
    for j, s in enumerate([s for s in sensors if s.kind == "FLOW" and hasattr(s, "_link")]):
        s.model_day = np.abs(_day_profile_from_sim(t, sim["Q"][:, j]))

    for s in sensors:
        if s.model_day is None or s.clean is None:
            continue
        x = s.clean[a:b].where(s.flag[a:b].isin([0, 1, 2, 3]))
        if s.kind == "PRESSURE" and np.isfinite(s.elev_meas):
            x = x + (s.elev_meas - s.elev)          # presión equivalente en la cota del nudo
        if x.notna().sum() < 48:
            continue
        meas_day = _typical_profile(x, getattr(s, 'onoff', False))
        s.meas_day = meas_day
        bias = float(np.nanmedian(x) - np.nanmean(s.model_day))
        amp_m = float(np.nanpercentile(x, 95) - np.nanpercentile(x, 5))
        amp_s = float(np.nanmax(s.model_day) - np.nanmin(s.model_day))
        r = float(np.corrcoef(meas_day, s.model_day)[0, 1]) if np.nanstd(s.model_day) > 1e-6 else np.nan
        s.cmp = {"sesgo": bias, "amp_med": amp_m, "amp_mod": amp_s, "r": r}
        if s.kind == "PRESSURE":
            hgl = np.nanmedian(x) + s.elev
            fe = net.zone_prv.get(s.zone, [])
            near_prv = None
            if s.xy is not None:
                for zs in net.zone_prv.values():
                    for v in zs:
                        d = float(np.hypot(*(net.xy[v["down"]] - np.array(s.xy))))
                        if d < 30:
                            near_prv = (v, d)
            if near_prv:
                v, d = near_prv
                uph = float(np.mean(up_head[v["up"]]))
                if s.zone == net.zone[v["down"]] and abs(hgl - uph) < abs(hgl - v["target"]) and hgl > v["target"] + 5:
                    s.note("error", f"Está a {d:.0f} m de la VRP {v['id']} y su carga medida ({hgl:.1f} m) se parece "
                                    f"más a la de aguas arriba ({uph:.1f} m) que a la consigna de aguas abajo "
                                    f"({v['target']:.1f} m): el equipo está casi seguro aguas arriba de la VRP. "
                                    f"Asígnelo al nudo {net.nid[v['up']]}.")
                    s.prv_side = "up"
            if fe and getattr(s, "prv_side", "") != "up":
                v = fe[0]
                uph = float(np.mean(up_head[v["up"]]))
                if hgl > v["target"] + 3 and amp_m > max(3 * amp_s, 5):
                    s.note("error", f"Carga medida {hgl:.1f} m, por encima de la consigna de la VRP "
                                    f"{v['id']} ({v['target']:.1f} m) y oscilando {amp_m:.0f} m al día "
                                    f"(el modelo {amp_s:.0f} m): la VRP no regulaba (abierta, by-pass o averiada) "
                                    f"y pasaba la presión de aguas arriba ({uph:.0f} m).")
                elif abs(hgl - uph) <= 3 and hgl > v["target"] + 5:
                    s.note("error", f"Carga medida {hgl:.1f} m igual a la de aguas arriba de la VRP {v['id']} "
                                    f"({uph:.1f} m) y {hgl - v['target']:.0f} m por encima de su consigna "
                                    f"({v['target']:.1f} m): la VRP no reduce (abierta, by-pass o fuera de servicio).")
                elif hgl > uph + 3:
                    s.note("error", f"Carga medida {hgl:.1f} m mayor que la de aguas arriba de su VRP "
                                    f"({uph:.1f} m): el sensor no está en la zona que dice el modelo "
                                    f"(ubicación o cota del nudo {s.element} incorrecta), o la VRP "
                                    f"{v['id']} está abierta y la zona se alimenta de otra.")
            s.hgl = hgl
        elif s.kind == "FLOW":
            ratio = np.nanmean(x) / max(np.nanmean(s.model_day), 1e-6)
            if abs(ratio - 1) > 0.25:
                s.note("aviso", f"Caudal medio medido {np.nanmean(x):.2f} L/s, modelo "
                                f"{np.nanmean(s.model_day):.2f} L/s (×{ratio:.2f}).")
            zero = (x.abs() < 0.05).mean()
            on = x[x.abs() > 0.05]
            if 0.15 < zero < 0.95 and len(on) > 10 and on.std() < 0.35 * on.mean():
                s.note("info", f"Funciona todo/nada ({zero * 100:.0f} % del tiempo a cero, "
                               f"{on.median():.2f} L/s cuando abre): válvula de boya o de altitud. "
                               f"Es real; para el día tipo vale el caudal medio.")
    _zone_bias(sensors, net)
    _similar_hints(sensors, window)
    return out


def _zone_bias(sensors, net):
    """Misma diferencia con el modelo en todos los sensores de una zona → consigna de la VRP;
    diferencias distintas → cota o pérdida de carga local."""
    byz = collections.defaultdict(list)
    for s in sensors:
        if s.kind == "PRESSURE" and getattr(s, "cmp", None) and s.zone is not None \
                and not any(lv == "error" for lv, _ in s.notes):
            byz[s.zone].append(s)
    for z, ss in byz.items():
        b = np.array([x.cmp["sesgo"] for x in ss])
        fe = net.zone_prv.get(z, [])
        if len(ss) >= 2 and abs(np.median(b)) > 2 and np.all(np.abs(b - np.median(b)) < 2.5) and fe:
            v = fe[0]
            txt = (f"Todos los sensores de la zona de la VRP {v['id']} ({', '.join(x.name for x in ss)}) están "
                   f"{np.median(b):+.1f} m respecto al modelo: la consigna real es ≈ {v['setting'] + np.median(b):.1f} m "
                   f"(en el INP {v['setting']:.1f} m). Corríjala antes de calibrar rugosidades.")
            for x in ss:
                x.note("aviso", txt)
            continue
        for x in ss:
            if abs(x.cmp["sesgo"]) > 5:
                others = [o for o in ss if o is not x]
                ctx = (f" (los demás de su zona: {', '.join(f'{o.name} {o.cmp['sesgo']:+.1f}' for o in others)})"
                       if others else "")
                x.note("aviso", f"Diferencia media con el modelo {x.cmp['sesgo']:+.1f} m{ctx}: cota del nudo "
                                f"{x.element} o del sensor, o pérdida de carga hasta este punto distinta de la del modelo.")


def _similar_hints(sensors, window):
    """Sensores con problemas de ubicación: ¿a quién se parecen?"""
    a, b = window
    ps = [s for s in sensors if s.kind == "PRESSURE" and s.clean is not None]
    prof = {}
    for s in ps:
        x = s.clean[a:b].where(s.flag[a:b].isin([0, 1, 2, 3]))
        pr = _typical_profile(x)
        if pr is not None:
            prof[s.name] = pr
    for s in ps:
        bad = s.element is None or any(lv == "error" and "VRP" in t and "aguas arriba de su" in t for lv, t in s.notes)
        if not bad or s.name not in prof:
            continue
        cands = []
        for o in ps:
            if o is s or o.name not in prof or o.element is None:
                continue
            r = np.corrcoef(prof[s.name], prof[o.name])[0, 1]
            dh = (np.nanmean(prof[s.name]) + (s.elev if s.element else 0)) - (np.nanmean(prof[o.name]) + o.elev)
            cands.append((r, o.name, dh, abs(np.nanmean(prof[s.name]) - np.nanmean(prof[o.name]))))
        cands.sort(reverse=True)
        top = [c for c in cands if c[0] > 0.8][:3]
        if not top:
            continue
        if s.element is None:
            s.note("info", "Por la forma de su día tipo se parece a " + ", ".join(
                f"{n} (r = {r:.2f}, presión media a {dp:.0f} m)" for r, n, _, dp in top) +
                ": probablemente está en esa zona; ubíquelo en el GPKG.")
        else:
            same = [c for c in top if abs(c[2]) < 6]
            if same:
                s.note("info", "Su carga y su forma coinciden con " + ", ".join(
                    f"{n} (r = {r:.2f}, Δcarga {dh:+.1f} m)" for r, n, dh, _ in same) +
                    ": en la realidad está en la misma zona de presión que ellos.")


def tank_balance(sensors: List[Sensor], net: Net, points: Dict[str, dict], log=print) -> List[str]:
    """Para cada tanque con caudalímetros de entrada y salida cerca: dV/dt = Qe − Qs."""
    msgs = []
    flows = [s for s in sensors if s.kind == "FLOW" and s.xy is not None]
    for i in np.where(net.ntype == en.EN_TANK)[0]:
        tid = net.nid[i]
        near = [s for s in flows if np.hypot(s.xy[0] - net.xy[i, 0], s.xy[1] - net.xy[i, 1]) < 60]
        def text(s):
            return (s.title + " " + points.get(s.point, {}).get("text", "")).lower()
        qin = [s for s in near if "entrada" in text(s)]
        qout = [s for s in near if "salida" in text(s)]
        if not (qin and qout):
            continue
        a = net.tank_diam.get(tid, 0)
        area = np.pi * a ** 2 / 4
        qi = sum(s.clean.where(s.flag != 9) for s in qin)
        qo = sum(s.clean.where(s.flag != 9) for s in qout)
        net_q = (qi - qo).dropna()
        if len(net_q) < 96 or area <= 0:
            continue
        daily = net_q.groupby(net_q.index.normalize()).mean() * 3.6     # m3/h
        rise = daily * 24 / area
        lo, hi = net.tank_lim.get(tid, (0, 0))
        names = f"{'+'.join(s.name for s in qin)} − {'+'.join(s.name for s in qout)}"
        if (rise.abs() > (hi - lo)).mean() > 0.5:
            sign = "entra" if daily.median() > 0 else "sale"
            m = (f"Tanque {tid} (Ø {a:.2f} m, {lo:.2f}–{hi:.2f} m): {names} = "
                 f"{daily.median():+.1f} m³/h de media ({net_q.median():+.2f} L/s, "
                 f"{abs(daily.median()) / max(qi.mean() * 3.6, 1e-6) * 100:.0f} % de la entrada). "
                 f"Equivale a {rise.median():+.1f} m/día, imposible en un tanque de {hi - lo:.1f} m: "
                 f"{'el tanque rebosa o tiene otra salida no medida (fuga, desagüe, toma)' if sign == 'entra' else 'hay otra entrada no medida'}, "
                 f"o un caudalímetro mide mal (≈{abs(net_q.median()):.2f} L/s de error).")
            msgs.append(m)
            for s in qin + qout:
                s.note("aviso", m)
        lvl = [s for s in sensors if s.kind == "LEVEL" and s.element == tid]
        for s in lvl:
            bad = s.flag == 9
            if bad.mean() > 0.3:
                # reconstrucción: integración del balance quitando la deriva diaria, centrada en
                # el nivel medio del modelo (o del rango del tanque)
                q = (qi - qo).reindex(s.clean.index).interpolate(limit=8)
                dq = q - q.rolling(96, center=True, min_periods=48).mean()
                h = (dq * 900 / 1000 / area).cumsum()
                h = h - h.rolling(96, center=True, min_periods=48).mean()
                mid = float(np.nanmean(s.model_day)) if s.model_day is not None else (lo + hi) / 2
                rec = (h + mid).clip(lo, hi)
                fill = bad & rec.notna() & (s.clean.index >= s.t_on) & (s.clean.index <= s.t_off)
                s.clean[fill] = rec[fill]
                s.flag[fill] = 4
                s.reason[fill] = "nivel reconstruido con el balance de caudales (forma; nivel medio del modelo)"
                s.note("info", "Nivel reconstruido con el balance de caudales de entrada y salida, sin la deriva "
                               "diaria y centrado en el nivel medio del modelo: sirve como forma de la curva, "
                               "no como nivel absoluto. Úselo solo como orientación o repita la medición.")
    return msgs


# --------------------------------------------------------------------------- veredicto y salidas
def verdicts(sensors: List[Sensor], window) -> None:
    a, b = window
    for s in sensors:
        if s.clean is None:
            s.verdict, s.weight = "No usar", 0
            continue
        f = s.flag[a:b]
        n = len(f) or 1
        pct = {k: (f == k).sum() / n * 100 for k in FLAG_TXT}
        s.pct = pct
        errors = [t for lv, t in s.notes if lv == "error"]
        if s.element is None:
            s.verdict, s.weight = "No usar (sin ubicación)", 0
        elif pct[9] > 50 or (s.kind == "LEVEL" and pct[4] > 30):
            s.verdict, s.weight = "No usar como meta", 0
        elif errors:
            s.verdict, s.weight = "Revisar antes de usar", 0.3
        elif pct[1] > 20 and float(s._noise[(s._noise.index >= a.normalize()) & (s._noise.index < b)].median()) \
                > 2 * CTX["noise_abs"][s.kind]:
            s.verdict, s.weight = "Usar solo la media horaria (tolerancia mayor)", 0.5
        elif pct[1] > 20:
            s.verdict, s.weight = "Válido (suavizado)", 0.9
        elif pct[2] + pct[3] > 20:
            s.verdict, s.weight = "Usar con precaución", 0.7
        else:
            s.verdict, s.weight = "Válido", 1.0


def write_dat(sensors: List[Sensor], window, out_dir: str, net: Net) -> Dict[str, str]:
    """Día tipo (mediana por cuarto de hora) del periodo de calibración, formato EPANET."""
    a, b = window
    files = {}
    for kind, fname in (("PRESSURE", "calibr_presion.dat"), ("FLOW", "calibr_caudal.dat"),
                        ("LEVEL", "calibr_nivel.dat")):
        lines = []
        for s in sensors:
            if s.kind != kind:
                continue
            if s.weight < 0.5 or s.element is None or getattr(s, "reservoir", False):
                lines.append(f";{s.name}: no se incluye ({s.verdict}"
                             f"{', embalse: su nivel va al patrón de carga' if getattr(s, 'reservoir', False) else ''})")
                continue
            x = s.clean[a:b].where(s.flag[a:b].isin([0, 1, 2, 3]))
            if kind == "PRESSURE" and np.isfinite(s.elev_meas):
                x = x + (s.elev_meas - s.elev)      # a la cota del nudo del INP
            prof = _typical_profile(x, getattr(s, 'onoff', False))
            if prof is None:
                continue
            if kind == "FLOW":
                prof = prof / net.q2lps     # a unidades del modelo
            lines.append(f";{s.name} ({s.verdict})")
            for k, v in enumerate(prof):
                lines.append(f"{s.element}\t{k * 0.25:.4f}\t{v:.4f}")
        p = os.path.join(out_dir, fname)
        if not any(not ln.startswith(";") for ln in lines):
            if os.path.exists(p):
                os.remove(p)
            continue
        if lines:
            with open(p, "w", encoding="utf-8") as fh:
                fh.write(f";Día tipo {fmt(a)} – {fmt(b)} (mediana por cuarto de hora), datos depurados"
                         f"{'; presiones pasadas a la cota del nudo con la cota GNSS' if kind == 'PRESSURE' else ''}\n")
                fh.write("\n".join(lines) + "\n")
            files[kind] = p
    return files


def write_excel(sensors, events, regimes, window, summary_rows, out_path) -> None:
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        pd.DataFrame(summary_rows).to_excel(xw, sheet_name="Resumen", index=False)
        ev = pd.DataFrame([{"inicio": e["inicio"], "fin": e["fin"], "tipo": e["tipo"],
                            "sensores": ", ".join(e["sensores"]), "interpretación": e["texto"]}
                           for e in events])
        ev.to_excel(xw, sheet_name="Eventos", index=False)
        pd.DataFrame([{"desde": r["desde"], "hasta": r["hasta"], "días": r["dias"],
                       "uso": r.get("uso", "")} for r in regimes]).to_excel(xw, sheet_name="Regímenes", index=False)
        a, b = window
        dia = {}
        for s in sensors:
            if s.clean is None:
                continue
            df = pd.DataFrame({"fecha": s.grid.index, "medido": s.grid.to_numpy(),
                               "corregido": (s.grid - s.offset).to_numpy(),
                               "depurado": s.clean.to_numpy(), "marca": s.flag.to_numpy(),
                               "método": [FLAG_TXT[int(k)] for k in s.flag], "motivo": s.reason.to_numpy()})
            df.to_excel(xw, sheet_name=s.name[:31], index=False)
            pr = _typical_profile(s.clean[a:b].where(s.flag[a:b].isin([0, 1, 2, 3])), getattr(s, 'onoff', False))
            if pr is not None:
                dia[f"{s.name} ({UNIT[s.kind]})"] = pr
        if dia:
            d = pd.DataFrame(dia)
            d.insert(0, "hora", [f"{k // 4:02d}:{(k % 4) * 15:02d}" for k in range(96)])
            d.to_excel(xw, sheet_name="Dia_tipo", index=False)
    _color_excel(out_path, [s.name[:31] for s in sensors if s.clean is not None])


def write_simple_excel(sensors, window, title, out_path) -> None:
    """Excel sencillo con formato corporativo: por sensor, dato inicial, marca, valor de
    reemplazo y dato depurado calculado con fórmula."""
    import openpyxl
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    RED, GREY, INK = "E30613", "F2F2F2", "1F1F1F"
    F = "Arial"
    head = Font(name=F, bold=True, color="FFFFFF", size=10)
    fill_head = PatternFill("solid", fgColor=RED)
    thin = Side(style="thin", color="D9D9D9")
    box = Border(bottom=thin)
    blue = Font(name=F, color="0000FF", size=10)
    normal = Font(name=F, color=INK, size=10)
    mark_fill = {1: "FFF2CC", 2: "DDEBF7", 3: "E2EFDA", 4: "EADCF8", 9: "F8CBAD"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name=F, bold=True, size=16, color=RED)
    ws["A2"] = (f"Periodo de calibración propuesto: {fmt(window[0])} – {fmt(window[1])}. "
                f"Dato depurado = fórmula sobre el dato inicial (ver leyenda).")
    ws["A2"].font = Font(name=F, size=10, color="595959")
    ws.row_dimensions[1].height = 26
    cols = ["Sensor", "Tipo", "Unidad", "Elemento INP", "Offset", "Uso recomendado", "Hoja"]
    for j, h in enumerate(cols, 1):
        c = ws.cell(4, j, h)
        c.font, c.fill = head, fill_head
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[4].height = 22
    usable = [s for s in sensors if s.clean is not None]
    names = {}
    for s in usable:
        nm = re.sub(r"[\[\]\*\?/\\:]", "_", s.name)[:31]
        names[s.name] = nm
    for i, s in enumerate(usable, 5):
        vals = [s.name, KIND_TXT[s.kind], UNIT[s.kind], s.element or "—", None, s.verdict, names[s.name]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font, c.border = normal, box
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GREY)
        ws.cell(i, 5).value = f"='{names[s.name]}'!$C$3"
        ws.cell(i, 5).number_format = "0.000"
        ws.cell(i, 7).hyperlink = f"#'{names[s.name]}'!A1"
        ws.cell(i, 7).font = Font(name=F, size=10, color=RED, underline="single")
    r = 6 + len(usable)
    ws.cell(r, 1, "Leyenda de la columna «Marca»").font = Font(name=F, bold=True, size=11, color=RED)
    legend = [(0, "válido", "Dato depurado = corregido = dato inicial − offset"),
              (1, "envolvente", "Oscilación más rápida que el registro: MEDIA.ACOTADA (TRIMMEAN, 50 %) del dato "
                                "corregido en la ventana de la hoja (±4 lecturas = 2 h, o ±8 = 4 h)"),
              (2, "interpolado", "Pico aislado o hueco ≤ 1 h: interpolación lineal (valor en «Reemplazo»)"),
              (3, "perfil típico", "Hueco de 1 a 6 h: perfil diario del propio sensor ajustado a los extremos "
                                   "(valor en «Reemplazo»)"),
              (4, "balance de masa", "Nivel de tanque reconstruido con los caudales de entrada y salida "
                                     "(valor en «Reemplazo»)"),
              (9, "excluido", "Sin instalar, corte de suministro, sensor fuera del agua: queda vacío")]
    for k, (m, n, t) in enumerate(legend, r + 1):
        c = ws.cell(k, 1, m)
        c.font = Font(name=F, bold=True, size=10)
        if m in mark_fill:
            c.fill = PatternFill("solid", fgColor=mark_fill[m])
        ws.cell(k, 2, n).font = Font(name=F, bold=True, size=10)
        ws.cell(k, 3, t).font = normal
    k = r + len(legend) + 2
    for t in ("En azul: valores editables (offset, marca y reemplazo). Si cambia una marca o el offset, "
              "el dato depurado se recalcula.",
              "Dato inicial: lectura del equipo llevada a la malla común de 15 min (interpolada en la hora real de "
              "cada lectura). Presión en m c.a., caudal en L/s, nivel en m."):
        ws.cell(k, 1, t).font = Font(name=F, size=9, italic=True, color="595959")
        k += 1
    for j, w in enumerate((14, 10, 8, 16, 10, 46, 14), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"

    for s in usable:
        sh = wb.create_sheet(names[s.name])
        sh.sheet_view.showGridLines = False
        sh["A1"] = f"{s.name} · {KIND_TXT[s.kind]} ({UNIT[s.kind]}) · {s.element or 'sin ubicación'}"
        sh["A1"].font = Font(name=F, bold=True, size=14, color=RED)
        sh["A2"] = s.verdict
        sh["A2"].font = Font(name=F, size=10, color="595959")
        off = s.offset if s.kind == "PRESSURE" else 0.0
        win = 8 if any("4 h" in str(x) for x in s.reason.unique()) else 4
        for cell, lab, val in (("B3", "Offset", off), ("B4", "Ventana (lecturas a cada lado)", win)):
            sh[cell.replace("B", "A")] = lab
            sh[cell.replace("B", "A")].font = Font(name=F, bold=True, size=10)
        sh["C3"], sh["C4"] = round(off, 4), win
        sh["C3"].font, sh["C4"].font = blue, normal
        sh["C3"].number_format = "0.000"
        hdr = ["Fecha y hora", f"Dato inicial ({UNIT[s.kind]})", "Marca", "Reemplazo",
               f"Corregido ({UNIT[s.kind]})", f"Dato depurado ({UNIT[s.kind]})", "Motivo"]
        R0 = 6
        for j, h in enumerate(hdr, 1):
            c = sh.cell(R0, j, h)
            c.font, c.fill = head, fill_head
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sh.row_dimensions[R0].height = 30
        g = s.grid
        n = len(g)
        first, last = R0 + 1, R0 + n
        for k in range(n):
            row = first + k
            fl = int(s.flag.iloc[k])
            v = g.iloc[k]
            sh.cell(row, 1, g.index[k].to_pydatetime()).number_format = "dd/mm/yyyy hh:mm"
            c = sh.cell(row, 2, None if pd.isna(v) else round(float(v), 4))
            c.number_format = "0.00"
            sh.cell(row, 3, fl).font = blue
            if fl in (2, 3, 4) and pd.notna(s.clean.iloc[k]):
                c = sh.cell(row, 4, round(float(s.clean.iloc[k]), 4))
                c.number_format, c.font = "0.00", blue
            a0, a1 = max(first, row - win), min(last, row + win)
            sh.cell(row, 5, f'=IF(OR(C{row}=9,B{row}=""),"",B{row}-$C$3)').number_format = "0.00"
            sh.cell(row, 6, f'=IF(C{row}=9,"",IF(C{row}=0,E{row},IF(C{row}=1,'
                            f'IF(COUNT(E{a0}:E{a1})=0,"",TRIMMEAN(E{a0}:E{a1},0.5)),IF(D{row}="","",D{row}))))')
            sh.cell(row, 6).number_format = "0.00"
            sh.cell(row, 6).font = Font(name=F, size=10, bold=True)
            if s.reason.iloc[k]:
                sh.cell(row, 7, s.reason.iloc[k])
        for row in sh.iter_rows(min_row=first, max_row=last, max_col=5):
            for c in row:
                if c.column in (1, 2, 5):
                    c.font = normal
        for m, colr in mark_fill.items():
            sh.conditional_formatting.add(f"C{first}:C{last}", FormulaRule(
                formula=[f"$C{first}={m}"], fill=PatternFill("solid", fgColor=colr)))
        for j, w in enumerate((18, 13, 8, 12, 13, 16, 58), 1):
            sh.column_dimensions[get_column_letter(j)].width = w
        sh.freeze_panes = sh.cell(first, 1)
        sh.auto_filter.ref = f"A{R0}:G{last}"
    wb.save(out_path)


def _color_excel(path, sheets):
    import openpyxl
    from openpyxl.styles import PatternFill
    fills = {1: "FFF2CC", 2: "DDEBF7", 3: "E2EFDA", 4: "EADCF8", 9: "F8CBAD"}
    wb = openpyxl.load_workbook(path)
    for sh in sheets:
        ws = wb[sh]
        for row in ws.iter_rows(min_row=2):
            k = row[4].value
            if k in fills:
                f = PatternFill("solid", fgColor=fills[k])
                for c in row[3:6]:
                    c.fill = f
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["F"].width = 30
        ws.column_dimensions["G"].width = 60
    for sh in ("Resumen", "Eventos"):
        ws = wb[sh]
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                90, max(10, max(len(str(c.value or "")) for c in col) * 0.9))
    wb.save(path)


# --------------------------------------------------------------------------- figuras e informe
COL = {0: "#1f2937", 1: "#d97706", 2: "#2563eb", 3: "#059669", 4: "#7c3aed"}


def _fig_sensor(s: Sensor, regimes, window, events) -> str:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.dates as md
    fig, ax = plt.subplots(1, 2, figsize=(13, 2.7), gridspec_kw={"width_ratios": [3.2, 1]})
    g = s.grid - (s.offset if s.kind == "PRESSURE" else 0)
    a0 = ax[0]
    a0.plot(g.index, g, color="#c0c4cc", lw=0.7, label="medido")
    for k in (0, 1, 2, 3, 4):
        m = s.flag == k
        if m.any():
            y = s.clean.where(m)
            a0.plot(y.index, y, color=COL[k], lw=1.0 if k else 0.9, label=FLAG_TXT[k])
    for i, j in runs((s.flag == 9).to_numpy()):
        a0.axvspan(s.flag.index[i], s.flag.index[min(j, len(s.flag) - 1)], color="#f87171", alpha=0.12, lw=0)
    a0.axvspan(window[0], window[1], color="#10b981", alpha=0.06, lw=0)
    for e in events:
        if e["tipo"] == "corte":
            a0.axvline(e["inicio"], color="#ef4444", lw=0.8, ls=":")
    for r in regimes[1:]:
        a0.axvline(r["desde"], color="#6b7280", lw=1, ls="--")
    ref = g[s.flag != 9] if (s.flag != 9).mean() > 0.3 else s.clean.dropna()
    lo, hi = np.nanpercentile(ref, [0.5, 99.5]) if len(ref) else (0, 1)
    pad = (hi - lo) * 0.15 + 0.5
    a0.set_ylim(lo - pad, hi + pad)
    a0.set_xlim(g.index[0], g.index[-1])
    a0.xaxis.set_major_locator(md.DayLocator())
    a0.xaxis.set_major_formatter(md.DateFormatter("%d/%m"))
    a0.set_ylabel(UNIT[s.kind])
    a0.grid(alpha=0.25)
    a0.legend(fontsize=7, ncol=6, loc="upper left", frameon=False)
    a1 = ax[1]
    h = np.arange(96) / 4
    if getattr(s, "meas_day", None) is not None:
        a1.plot(h, s.meas_day, color="#E30613", lw=1.6, label="día tipo depurado")
    if s.model_day is not None and not getattr(s, "reservoir", False):
        a1.plot(h, s.model_day, color="#374151", lw=1.2, ls="--", label="modelo")
    a1.set_xlim(0, 24)
    a1.set_xticks([0, 6, 12, 18, 24])
    a1.grid(alpha=0.25)
    a1.legend(fontsize=7, frameon=False)
    a1.set_title("día tipo", fontsize=8)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=80)
    plt.close(fig)
    return base64.b64encode(buf.getvalue()).decode()


def write_report(path, title, story, sensors, events, regimes, window, balance_msgs, files):
    esc = html.escape
    lv_icon = {"error": ("✖", "#b91c1c"), "aviso": ("▲", "#b45309"), "info": ("ℹ", "#1d4ed8")}
    rows = []
    for s in sensors:
        pct = getattr(s, "pct", {})
        rows.append(
            f"<tr><td><b>{esc(s.name)}</b></td><td>{KIND_TXT[s.kind]}</td><td>{esc(str(s.element or '—'))}</td>"
            f"<td>{esc(str(s.prv or ''))}</td>"
            f"<td>{pct.get(0, 0):.0f}</td><td>{pct.get(1, 0):.0f}</td><td>{pct.get(2, 0) + pct.get(3, 0):.0f}</td>"
            f"<td>{pct.get(4, 0):.0f}</td><td>{pct.get(9, 0):.0f}</td>"
            f"<td class='v{int(s.weight * 10)}'>{esc(s.verdict)}</td></tr>")
    cards = []
    for s in sensors:
        if s.clean is None:
            notes = "".join(f"<li style='color:{lv_icon[l][1]}'>{lv_icon[l][0]} {esc(t)}</li>" for l, t in s.notes)
            cards.append(f"<div class='card'><h3>{esc(s.name)}</h3><ul>{notes}</ul></div>")
            continue
        img = _fig_sensor(s, regimes, window, events)
        notes = "".join(f"<li><span style='color:{lv_icon[l][1]}'>{lv_icon[l][0]}</span> {esc(t)}</li>"
                        for l, t in s.notes)
        cmp_ = getattr(s, "cmp", None)
        cm = (f"<p class='m'>Frente al modelo (periodo de calibración): diferencia media {cmp_['sesgo']:+.2f} "
              f"{UNIT[s.kind]}, oscilación diaria medida {cmp_['amp_med']:.1f} / modelo {cmp_['amp_mod']:.1f}, "
              f"r = {cmp_['r']:.2f}</p>") if cmp_ else ""
        cards.append(
            f"<div class='card'><h3>{esc(s.name)} <small>{KIND_TXT[s.kind]} · {esc(str(s.element))}"
            f"{' · zona de la VRP ' + esc(s.prv) if s.prv else ''}</small>"
            f"<span class='badge v{int(s.weight * 10)}'>{esc(s.verdict)}</span></h3>"
            f"<img src='data:image/png;base64,{img}'/>{cm}<ul>{notes}</ul></div>")
    ev = "".join(f"<tr><td>{fmt(e['inicio'])}</td><td>{fmt(e['fin'])}</td><td>{esc(e['tipo'])}</td>"
                 f"<td>{esc(e['texto'])}</td></tr>" for e in events)
    rg = "".join(f"<tr><td>{fmt(r['desde'])}</td><td>{fmt(r['hasta'])}</td><td>{r['dias']}</td>"
                 f"<td>{esc(r.get('uso', ''))}</td></tr>" for r in regimes)
    st = "".join(f"<li>{s}</li>" for s in story)
    fl = "".join(f"<li><code>{esc(os.path.basename(v))}</code></li>" for v in files)
    doc = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{esc(title)}</title>
<style>
:root{{--ink:#111827;--mut:#6b7280;--bg:#f3f4f6;--card:#fff;--acc:#E30613}}
body{{font-family:Segoe UI,system-ui,sans-serif;background:var(--bg);color:var(--ink);margin:0;padding:24px;line-height:1.45}}
h1{{margin:0 0 4px}} h2{{margin-top:28px;border-left:4px solid var(--acc);padding-left:10px}}
.card{{background:var(--card);border-radius:12px;padding:14px 18px;margin:12px 0;box-shadow:0 1px 2px #0001}}
.card img{{width:100%;max-width:1040px}}
table{{border-collapse:collapse;background:#fff;border-radius:10px;overflow:hidden;font-size:13px}}
td,th{{padding:6px 10px;border-bottom:1px solid #e5e7eb;text-align:left;vertical-align:top}} th{{background:#f9fafb}}
small{{color:var(--mut);font-weight:400}} .m{{color:var(--mut);font-size:13px;margin:4px 0}}
.badge{{float:right;font-size:12px;padding:3px 10px;border-radius:99px;font-weight:600}}
.v10{{background:#dcfce7;color:#166534}} .v7{{background:#e0f2fe;color:#075985}} .v5{{background:#fef3c7;color:#92400e}}
.v3{{background:#ffedd5;color:#9a3412}} .v0{{background:#fee2e2;color:#991b1b}}
.leg span{{display:inline-block;margin-right:14px;font-size:13px}} .leg i{{display:inline-block;width:18px;height:3px;vertical-align:middle;margin-right:5px}}
li{{margin:3px 0}}
</style></head><body>
<h1>{esc(title)}</h1>
<p class="m">Periodo propuesto para calibrar: <b>{fmt(window[0])} – {fmt(window[1])}</b> (franja verde en las figuras).</p>
<h2>Qué pasó probablemente en la campaña</h2><div class="card"><ol>{st}</ol></div>
<h2>Regímenes de operación</h2><table><tr><th>desde</th><th>hasta</th><th>días</th><th>uso</th></tr>{rg}</table>
<h2>Eventos</h2><table><tr><th>inicio</th><th>fin</th><th>tipo</th><th>interpretación</th></tr>{ev}</table>
<h2>Qué datos tomar y cuáles reemplazar (periodo de calibración, % de los datos)</h2>
<table><tr><th>sensor</th><th>tipo</th><th>elemento INP</th><th>VRP de su zona</th><th>válido</th><th>envolvente</th>
<th>interpolado / perfil</th><th>balance</th><th>excluido</th><th>uso</th></tr>{''.join(rows)}</table>
<p class="leg"><span><i style="background:#1f2937"></i>válido</span><span><i style="background:#d97706"></i>envolvente (media recortada 2 h)</span>
<span><i style="background:#2563eb"></i>interpolado</span><span><i style="background:#059669"></i>perfil típico</span>
<span><i style="background:#7c3aed"></i>balance de masa</span><span><i style="background:#f87171"></i>excluido (sombreado)</span>
<span>-- cambio de régimen · ⋯ corte</span></p>
<h2>Sensores</h2>{''.join(cards)}
<h2>Ficheros</h2><ul>{fl}</ul>
</body></html>"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(doc)


# --------------------------------------------------------------------------- relato
def build_story(sensors, events, regimes, window, balance_msgs, net) -> List[str]:
    S = []
    esc = html.escape
    # régimen elegido
    if len(regimes) > 1:
        lines = []
        for r in regimes:
            lines.append(f"{fmt(r['desde'])} – {fmt(r['hasta'])} ({r['dias']} d)")
        S.append("La campaña tiene <b>" + str(len(regimes)) + " regímenes de operación distintos</b>: "
                 + "; ".join(lines) + ". Un modelo representa una sola forma de operar la red: no se deben "
                 "mezclar en la calibración.")
    timed = []
    empt = [e for e in events if e["tipo"] == "tanque vacío"]
    for e in empt:
        after = [f for f in events if f["tipo"] == "caudal nulo"
                 and e["inicio"] - pd.Timedelta("3h") <= f["inicio"] <= e["fin"] + pd.Timedelta("6h")]
        dry = [f for f in events if f["tipo"] in ("sin presión", "corte")
               and e["inicio"] - pd.Timedelta("2h") <= f["inicio"] <= e["fin"]]
        t = (f"<b>{fmt(e['inicio'])}–{fmt(e['fin'])}</b>: {esc(e['texto'])}")
        if after:
            t += (" Al mismo tiempo " + ", ".join(f"{f['sensores'][0]} cae a casi cero "
                                                  f"({fmt(f['inicio'])}–{fmt(f['fin'])})" for f in after) +
                  ": se cerró o se estranguló la salida para recuperar el tanque (racionamiento).")
        if dry:
            t += " " + ", ".join(sorted({f["sensores"][0] for f in dry})) + " se queda sin presión mientras el tanque está vacío."
        timed.append((e["inicio"], t))
    cuts = [e for e in events if e["tipo"] == "corte"]
    for e in cuts:
        nxt = [r for r in regimes if e["fin"] <= r["desde"] <= e["fin"] + pd.Timedelta("1D")]
        extra = (f" A las {fmt(nxt[0]['desde'])} la red cambia de forma de operar (ver más abajo): la maniobra "
                 f"de este corte es probablemente el origen del cambio.") if nxt else ""
        timed.append((e["inicio"], f"<b>{fmt(e['inicio'])}</b>: {esc(e['texto'])} Duración "
                                   f"{(e['fin'] - e['inicio']).total_seconds() / 3600:.1f} h.{extra}"))
    S.extend(t for _, t in sorted(timed, key=lambda x: x[0]))
    # VRP que no regulan
    prv_bad = collections.defaultdict(list)
    for s in sensors:
        for lv, t in s.notes:
            if "no regulaba" in t:
                prv_bad[s.prv].append(s.name)
    # comparar regímenes en sensores con VRP
    for prv, who in prv_bad.items():
        S.append(f"La <b>VRP {esc(str(prv))}</b> no regulaba en el periodo de calibración: {', '.join(who)} "
                 f"siguen la presión de aguas arriba.")
    reg_txt = _regime_contrast(sensors, regimes, net)
    S.extend(reg_txt)
    fast = collections.defaultdict(list)
    for s in sensors:
        for d in getattr(s, "_fast_days", []):
            fast[s.name].append(d)
    if fast:
        allx = sorted({d for v in fast.values() for d in v})
        S.append("Oscilación rápida (más rápida que el registro de 15 min) en <b>"
                 + ", ".join(f"{k} ({len(v)} d)" for k, v in fast.items()) + "</b>, desde el "
                 + pd.Timestamp(allx[0]).strftime("%d/%m") + ". El componente rápido no está correlacionado "
                 "entre sensores: cada equipo muestrea la oscilación en un segundo distinto, así que el valor "
                 "instantáneo no es recuperable. Suele ser una VRP que bombea (sobredimensionada a caudal bajo) "
                 "o aire en la red tras un vaciado. Se toma la media recortada de 2 h (4 h si la oscilación es muy fuerte), que es lo que reproduce un modelo "
                 "en periodo extendido.")
    for m in balance_msgs:
        S.append(esc(m))
    for s in sensors:
        for lv, t in s.notes:
            if lv == "error" and "no regulaba" not in t:
                S.append(f"<b>{esc(s.name)}</b>: {esc(t)}")
    return S


def _regime_contrast(sensors, regimes, net) -> List[str]:
    if len(regimes) < 2:
        return []
    out = []
    for s in sensors:
        if s.kind != "PRESSURE" or s.clean is None or not s.prv:
            continue
        st = []
        for r in regimes:
            x = s.clean[r["desde"]:r["hasta"]].where(s.flag[r["desde"]:r["hasta"]].isin([0, 1, 2, 3])).dropna()
            if len(x) < 48:
                st.append(None)
                continue
            st.append((float(x.median() + s.elev), float(x.quantile(0.95) - x.quantile(0.05))))
        ok = [v for v in st if v]
        if len(ok) < 2:
            continue
        v = net.zone_prv[s.zone][0]
        amps = [a for _, a in ok]
        if max(amps) > 3 * max(min(amps), 3) and max(amps) > 15:
            i_hi = int(np.argmax([a if a is not None else -1 for a in [x[1] if x else None for x in st]]))
            i_lo = int(np.argmin([a if a is not None else 1e9 for a in [x[1] if x else None for x in st]]))
            out.append(f"<b>{s.name}</b> (zona de la VRP {v['id']}, consigna → carga {v['target']:.0f} m): "
                       f"oscila {st[i_hi][1]:.0f} m al día con carga media {st[i_hi][0]:.0f} m en "
                       f"{fmt(regimes[i_hi]['desde'])}–{fmt(regimes[i_hi]['hasta'])}, y {st[i_lo][1]:.0f} m con "
                       f"{st[i_lo][0]:.0f} m en {fmt(regimes[i_lo]['desde'])}–{fmt(regimes[i_lo]['hasta'])}. "
                       f"En el primero la VRP no regulaba (by-pass o VRP abierta); en el segundo sí.")
    return out


# --------------------------------------------------------------------------- flujo principal
CTX = {"noise_floor": {"PRESSURE": 0.5, "LEVEL": 0.03, "FLOW": 0.2},
       "noise_abs": {"PRESSURE": 3.0, "LEVEL": 0.15, "FLOW": 1.0},
       "spike_abs": {"PRESSURE": 8.0, "LEVEL": 0.3, "FLOW": 2.0}}


def run(inp: str, excel: str, gpkgs: Sequence[str], out_dir: str, cfg: Optional[dict] = None,
        lib: Optional[str] = None, log: Callable[[str], None] = print) -> dict:
    cfg = cfg or {}
    os.makedirs(out_dir, exist_ok=True)
    log("Leyendo la campaña…")
    sensors = read_campaign_excel(excel, log, cfg.get("sheets"))
    points = {}
    for g in gpkgs:
        points.update(read_gpkg_points(g))
    log(f"Puntos en GeoPackage: {', '.join(points)}")
    net = Net(inp, lib)
    for s in sensors:
        s.grid = to_grid(s.raw)
        detect_service(s)
        analyse_sensor(s, CTX)
    log("Ubicando los sensores en el modelo…")
    locate(sensors, points, net, log)
    events = find_events(sensors)
    on = pd.Series([s.t_on for s in sensors if s.kind == "PRESSURE" and s.t_on is not None])
    if len(on):
        t0 = on.quantile(0.2) - pd.Timedelta("1D")
        events = [e for e in events if e["fin"] >= t0]
    regimes = find_regimes(sensors)
    if not regimes:
        t0 = max(s.t_on for s in sensors if s.t_on is not None)
        t1 = min(s.t_off for s in sensors if s.t_off is not None)
        regimes = [{"desde": t0, "hasta": t1, "dias": int((t1 - t0).days)}]
    # periodo de calibración: el régimen estable más largo dentro del servicio de los equipos
    if "window" in cfg:
        window = (pd.Timestamp(cfg["window"][0]), pd.Timestamp(cfg["window"][1]))
    else:
        t_on = pd.Series([s.t_on for s in sensors if s.kind == "PRESSURE" and s.t_on is not None]).quantile(0.8)
        t_off = pd.Series([s.t_off for s in sensors if s.kind == "PRESSURE" and s.t_off is not None]).quantile(0.2)
        best = None
        for r in regimes:
            a, b = max(r["desde"], t_on), min(r["hasta"], t_off)
            # evitar las horas posteriores a un corte
            for e in events:
                if e["tipo"] == "corte" and a <= e["fin"] <= b and (e["fin"] - a) < pd.Timedelta("1D"):
                    a = e["fin"] + pd.Timedelta("2h")
            if best is None or (b - a) > (best[1] - best[0]):
                best = (a.ceil("1h"), b.floor("1h"))
        window = best
    for r in regimes:
        r["uso"] = "CALIBRACIÓN" if r["desde"] <= window[0] < r["hasta"] else "no mezclar con la calibración"
    log(f"Periodo de calibración: {fmt(window[0])} – {fmt(window[1])}")

    log("Comparando con el modelo…")
    model_checks(sensors, net, window, log)
    balance_msgs = tank_balance(sensors, net, points, log)
    # tras reconstruir niveles, recomparar el perfil
    verdicts(sensors, window)
    story = build_story(sensors, events, regimes, window, balance_msgs, net)

    summary = []
    for s in sensors:
        pct = getattr(s, "pct", {})
        summary.append({"sensor": s.name, "tipo": KIND_TXT[s.kind], "punto GPKG": s.point,
                        "elemento INP": s.element, "cota modelo": None if np.isnan(s.elev) else round(s.elev, 2),
                        "VRP de su zona": s.prv, "instalado": s.t_on, "retirado": s.t_off,
                        "offset (m)": round(s.offset, 3),
                        **{f"% {FLAG_TXT[k]}": round(pct.get(k, 0), 1) for k in FLAG_TXT},
                        "uso": s.verdict, "peso sugerido": s.weight,
                        "diagnóstico": " | ".join(t for _, t in s.notes)})
    files = []
    xl = os.path.join(out_dir, "campaña_depurada.xlsx")
    write_excel(sensors, events, regimes, window, summary, xl)
    files.append(xl)
    simple = os.path.join(out_dir, "datos_depurados.xlsx")
    write_simple_excel(sensors, window, f"Datos depurados · {os.path.splitext(os.path.basename(excel))[0]}", simple)
    files.append(simple)
    dats = write_dat(sensors, window, out_dir, net)
    files += list(dats.values())
    rep = os.path.join(out_dir, "informe_campaña.html")
    files.append(rep)
    write_report(rep, f"Análisis de la campaña · {os.path.basename(excel)}", story, sensors, events,
                 regimes, window, balance_msgs, files)
    log(f"Informe: {rep}")
    return {"sensors": sensors, "events": events, "regimes": regimes, "window": window,
            "story": story, "files": files, "dat": dats}


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(prog="python -m epacal.campaign",
                                 description="Analiza y depura una campaña de medición con el INP.")
    ap.add_argument("--inp", required=True)
    ap.add_argument("--excel", required=True)
    ap.add_argument("--gpkg", nargs="+", default=[])
    ap.add_argument("--out", default="campaña")
    ap.add_argument("--config", help="YAML opcional: sheets:{hoja:{type,col,name,point,skip}}, window:[ini, fin]")
    a = ap.parse_args(argv)
    cfg = {}
    if a.config:
        import yaml
        with open(a.config, encoding="utf-8") as fh:
            cfg = yaml.safe_load(fh) or {}
    run(a.inp, a.excel, a.gpkg, a.out, cfg)


if __name__ == "__main__":
    main()
